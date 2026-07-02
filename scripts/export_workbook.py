# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl"]
# ///
"""Export the events JSON produced by a scrape run into the output workbook.

Usage: uv run scripts/export_workbook.py data/events_YYYYMMDD.json

Applies the V1 subset of QC rules from the constitution (Section 11) and
fails loudly if any check is violated. Never overwrites an existing workbook.
"""

import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook

# Section 5 taxonomy: category -> valid subcategory prefixes
TAXONOMY = {
    "A_AGE_ELIGIBILITY_ASSURANCE": ["A1_MINIMUM_AGE_POLICY", "A2_AGE_PREDICTION_OR_AGE_INFERENCE", "A3_ID_CHECK_OR_AGE_VERIFICATION", "A4_DEFAULT_TEEN_EXPERIENCE", "A5_UNDER_13_SUPERVISED_ACCESS", "A6_AGE_ENFORCEMENT_AND_ACCOUNT_ACTIONS"],
    "B_TEEN_MODE_YOUTH_SAFETY_DEFAULTS": ["B1_STRICTER_CONTENT_FILTERS_FOR_MINORS", "B2_REDUCED_PERSONALIZATION_OR_MEMORY_FOR_MINORS", "B3_SENSITIVE_TOPIC_HANDLING_FOR_TEENS", "B4_EDUCATION_OR_SCHOOL_ACCOUNT_PROTECTIONS", "B5_YOUTH_DEFAULT_PRIVACY_SETTINGS", "B6_SAFE_COMPLETION_STYLE_FOR_MINORS"],
    "C_PARENTAL_FAMILY_CONTROLS": ["C1_PARENT_DASHBOARD_OR_FAMILY_CENTER", "C2_DISABLE_OR_RESTRICT_AI_FEATURES_FOR_MINOR", "C3_PARENTAL_VISIBILITY_AND_MINOR_PRIVACY_BALANCE", "C4_CRISIS_OR_SELF_HARM_ALERTS_TO_PARENT_TRUSTED_CONTACT", "C5_PARENTAL_EDUCATION_AND_GUIDANCE"],
    "D_SELF_HARM_CRISIS_HANDLING": ["D1_988_OR_CRISIS_RESOURCE_ROUTING", "D2_SELF_HARM_OR_SUICIDE_RESPONSE_PROTOCOL", "D3_EMERGENCY_ESCALATION_OR_DUTY_TO_ACT_POLICY", "D4_TRUSTED_ADULT_OR_CLINICIAN_DISCLOSURE_PROMPT", "D5_EATING_DISORDER_SUBSTANCE_ABUSE_OR_ABUSE_CRISIS_HANDLING", "D6_NOT_A_CRISIS_SERVICE_DISCLAIMER"],
    "E_HARMFUL_CONTENT_REFUSAL_SAFETY_FILTERS": ["E1_SELF_HARM_METHODS_OR_INSTRUCTIONS_REFUSAL", "E2_VIOLENCE_ABUSE_EXPLOITATION_REFUSAL", "E3_SEXUAL_CONTENT_INVOLVING_MINORS_REFUSAL", "E4_ILLICIT_SUBSTANCES_OR_MEDICAL_MISUSE_REFUSAL", "E5_EATING_DISORDER_OR_BODY_HARM_REFUSAL", "E6_EVASION_JAILBREAK_OR_UNSAFE_COACHING_REFUSAL", "E7_BULLYING_HARASSMENT_OR_COERCION_REFUSAL"],
    "F_MENTAL_HEALTH_ADVICE_BOUNDARIES": ["F1_DIAGNOSIS_TREATMENT_REFUSAL_OR_LIMITATION", "F2_NOT_THERAPIST_OR_NOT_MEDICAL_ADVICE_DISCLOSURE", "F3_MEDICATION_OR_CLINICAL_CARE_ADVICE_LIMITS", "F4_SUPPORTIVE_NONCLINICAL_FRAMING", "F5_THERAPIST_ROLEPLAY_OR_SIMULATED_CLINICIAN_RESTRICTION", "F6_PROFESSIONAL_HELP_SEEKING_PROMPTS"],
    "G_PRIVACY_MEMORY_RETENTION_TRAINING": ["G1_CHAT_HISTORY_RETENTION_POLICY", "G2_MEMORY_CONTROLS", "G3_MODEL_TRAINING_OPT_OUT_OR_DATA_USE", "G4_HUMAN_REVIEW_OF_CONVERSATIONS", "G5_SENSITIVE_CONVERSATION_HANDLING", "G6_THIRD_PARTY_SHARING_OR_SUBPROCESSORS", "G7_DELETION_EXPORT_OR_ACCOUNT_DATA_RIGHTS"],
    "H_TRANSPARENCY_REPORTING_EVALUATION": ["H1_RELEASE_NOTES_OR_CHANGELOG", "H2_MODEL_CARD_OR_SYSTEM_CARD", "H3_SAFETY_REPORT_OR_TRANSPARENCY_REPORT", "H4_RED_TEAMING_OR_INDEPENDENT_EVALUATION", "H5_POLICY_VERSIONING_OR_DOCUMENTED_GOVERNANCE_PROCESS"],
    "I_PRODUCT_MODEL_INTEGRATION_EVENTS": ["I1_NEW_MODEL_RELEASE_OR_MODEL_REPLACEMENT", "I2_VOICE_VIDEO_OR_MULTIMODAL_ROLLOUT", "I3_COMPANION_CHARACTER_OR_ROLEPLAY_MODE", "I4_SOCIAL_MEDIA_OR_MESSAGING_INTEGRATION", "I5_API_OR_DEVELOPER_POLICY_CHANGE", "I6_HEALTHCARE_EDUCATION_OR_ENTERPRISE_INTEGRATION"],
    "J_ACCESS_MARKET_APP_STORE_ACCOUNT": ["J1_APP_STORE_AGE_RATING_OR_METADATA", "J2_GEOGRAPHIC_ROLLOUT_OR_RESTRICTION", "J3_ACCOUNT_ELIGIBILITY_OR_CONSENT_RULES", "J4_FEATURE_ACCESS_RESTRICTION_OR_PERMISSIONING"],
    "K_INCIDENT_RESPONSE_PRODUCT_CHANGE": ["K1_SAFETY_INCIDENT_RESPONSE", "K2_LITIGATION_SETTLEMENT_OR_REGULATORY_RESPONSE", "K3_PUBLIC_CONTROVERSY_GOVERNANCE_CHANGE"],
}

EVENT_FIELDS = [
    "event_id", "company", "platform", "event_title", "event_summary",
    "event_importance", "event_type", "event_date", "date_basis",
    "date_confidence", "primary_category_code", "primary_subcategory_code",
    "safety_relevance_score", "expected_direction", "affected_user_group",
    "analytic_priority", "mechanism_summary", "survey_exposure_mapping",
    "source_name", "source_url", "supporting_source_urls", "relevant_quote",
    "classification_confidence", "human_review_needed", "human_review_reason",
]

RUN_LOG_FIELDS = [
    "source_name", "source_url", "company_hint", "platform_hint",
    "status", "records_found", "records_included", "error_message",
]

REVIEW_FIELDS = [
    "event_id", "review_reason", "suggested_question_for_reviewer",
    "priority", "source_url",
]

ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")

IMPORTANCE_VALUES = {"major_safety_relevant_event", "minor_contextual_update", "baseline_policy_snapshot", "exclude"}
CONFIDENCE_VALUES = {"high", "medium", "low"}


def qc(events: list[dict]) -> list[str]:
    errors = []
    seen_ids = set()
    for ev in events:
        eid = ev.get("event_id", "<missing event_id>")
        if eid in seen_ids:
            errors.append(f"{eid}: duplicate event_id")
        seen_ids.add(eid)
        for field in EVENT_FIELDS:
            if field in ("supporting_source_urls", "human_review_reason"):
                continue
            if ev.get(field) in (None, ""):
                errors.append(f"{eid}: missing required field {field}")
        cat = ev.get("primary_category_code")
        sub = ev.get("primary_subcategory_code")
        if cat not in TAXONOMY:
            errors.append(f"{eid}: unknown primary_category_code {cat}")
        elif sub not in TAXONOMY[cat]:
            errors.append(f"{eid}: subcategory {sub} not valid under {cat}")
        date = ev.get("event_date", "")
        if date != "unknown" and not ISO_DATE.match(str(date)):
            errors.append(f"{eid}: event_date {date!r} is neither ISO YYYY-MM-DD nor 'unknown'")
        if ev.get("event_importance") not in IMPORTANCE_VALUES:
            errors.append(f"{eid}: invalid event_importance {ev.get('event_importance')}")
        if ev.get("classification_confidence") not in CONFIDENCE_VALUES:
            errors.append(f"{eid}: invalid classification_confidence")
        if ev.get("classification_confidence") == "low" and not ev.get("human_review_needed"):
            errors.append(f"{eid}: low classification_confidence requires human_review_needed")
        if ev.get("event_importance") == "major_safety_relevant_event" and int(ev.get("safety_relevance_score", 0)) < 2:
            errors.append(f"{eid}: major event requires safety_relevance_score >= 2")
        if ev.get("human_review_needed") and not ev.get("human_review_reason"):
            errors.append(f"{eid}: human_review_needed without human_review_reason")
        if not str(ev.get("source_url", "")).startswith("http"):
            errors.append(f"{eid}: source_url is not a URL")
    return errors


def main() -> None:
    if len(sys.argv) != 2:
        sys.exit("Usage: uv run scripts/export_workbook.py data/events_YYYYMMDD.json")
    data_path = Path(sys.argv[1])
    data = json.loads(data_path.read_text())
    events = data["events"]
    run_log = data["source_run_log"]
    run_date = data["run_date"].replace("-", "")

    errors = qc(events)
    if errors:
        print("QC FAILED:")
        for err in errors:
            print(f"  - {err}")
        sys.exit(1)

    out_path = Path(__file__).resolve().parent.parent / "output" / f"platform_governance_events_output_{run_date}.xlsx"
    if out_path.exists():
        sys.exit(f"Refusing to overwrite existing workbook: {out_path}")

    wb = Workbook()

    ws = wb.active
    ws.title = "platform_events"
    ws.append(EVENT_FIELDS + ["run_id"])
    for ev in events:
        row = [ev.get(f, "") for f in EVENT_FIELDS] + [data["run_id"]]
        ws.append(["TRUE" if v is True else "FALSE" if v is False else v for v in row])

    ws_log = wb.create_sheet("source_run_log")
    ws_log.append(RUN_LOG_FIELDS + ["run_id"])
    for src in run_log:
        ws_log.append([src.get(f, "") for f in RUN_LOG_FIELDS] + [data["run_id"]])

    ws_rev = wb.create_sheet("review_queue")
    ws_rev.append(REVIEW_FIELDS)
    review_rows = 0
    for ev in events:
        if ev.get("human_review_needed"):
            ws_rev.append([
                ev["event_id"],
                ev.get("human_review_reason", ""),
                ev.get("suggested_question_for_reviewer", ""),
                ev.get("review_priority", "medium"),
                ev.get("source_url", ""),
            ])
            review_rows += 1

    wb.save(out_path)
    print(f"QC passed ({len(events)} events checked)")
    print(f"Wrote {out_path}")
    print(f"  platform_events: {len(events)} rows")
    print(f"  source_run_log: {len(run_log)} rows")
    print(f"  review_queue: {review_rows} rows")


if __name__ == "__main__":
    main()
